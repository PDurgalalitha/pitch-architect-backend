import io
import base64
import requests
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

from app.schemas.pitch_deck import PitchDeckResponse

def hex_to_rgb(hex_str: str) -> RGBColor:
    hex_str = hex_str.lstrip('#')
    if len(hex_str) != 6:
        return RGBColor(15, 23, 42)
    return RGBColor(*(int(hex_str[i:i+2], 16) for i in (0, 2, 4)))

def generate_pptx_from_deck(deck: PitchDeckResponse) -> bytes:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    for index, slide_data in enumerate(deck.slides, start=1):
        slide = prs.slides.add_slide(blank_layout)
        
        font_name = slide_data.font_family or "Arial"
        color_text = hex_to_rgb(slide_data.text_color)
        color_bg = hex_to_rgb(slide_data.bg_color)
        color_accent = hex_to_rgb(slide_data.accent_color)
        
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color_bg

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.733), Inches(0.5))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = f"SLIDE {index}: {slide_data.slide_title}".upper()
        p_title.font.name = font_name
        p_title.font.size = Pt(12)
        p_title.font.bold = True
        p_title.font.color.rgb = color_accent

        headline_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.0), Inches(11.733), Inches(0.8))
        tf_headline = headline_box.text_frame
        tf_headline.word_wrap = True
        p_head = tf_headline.paragraphs[0]
        p_head.text = slide_data.headline
        p_head.font.name = font_name
        p_head.font.size = Pt(slide_data.headline_size or 24)
        p_head.font.bold = True
        p_head.font.color.rgb = color_text

        has_image = bool(slide_data.image_url)
        content_width = Inches(6.2) if has_image else Inches(11.733)
        bullets_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.9), content_width, Inches(4.2))
        tf_bullets = bullets_box.text_frame
        tf_bullets.word_wrap = True

        for idx, bullet in enumerate(slide_data.bullet_points):
            p_b = tf_bullets.add_paragraph() if idx > 0 else tf_bullets.paragraphs[0]
            p_b.text = f"• {bullet}"
            p_b.font.name = font_name
            p_b.font.size = Pt(slide_data.body_size or 14)
            p_b.font.color.rgb = color_text
            p_b.space_after = Pt(10)

        if has_image:
            try:
                raw_bytes = None
                if slide_data.image_url.startswith("data:image/"):
                    image_data = slide_data.image_url.split(",", 1)[1]
                    raw_bytes = base64.b64decode(image_data)
                else:
                    resp = requests.get(slide_data.image_url, timeout=10)
                    if resp.status_code == 200:
                        raw_bytes = resp.content
                
                if raw_bytes:
                    img = Image.open(io.BytesIO(raw_bytes))
                    png_stream = io.BytesIO()
                    img.save(png_stream, format="PNG")
                    png_stream.seek(0)
                    slide.shapes.add_picture(png_stream, Inches(7.3), Inches(1.9), Inches(5.2), Inches(4.2))
            except Exception as e:
                print(f"Error rendering PPTX image for slide {index}: {e}")

        takeaway_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.3), Inches(11.733), Inches(0.8))
        tf_takeaway = takeaway_box.text_frame
        tf_takeaway.word_wrap = True
        p_take = tf_takeaway.paragraphs[0]
        p_take.text = f"Core Takeaway: {slide_data.key_takeaway}"
        p_take.font.name = font_name
        p_take.font.size = Pt(12)
        p_take.font.bold = True
        p_take.font.color.rgb = color_accent

    output_stream = io.BytesIO()
    prs.save(output_stream)
    output_stream.seek(0)
    return output_stream.getvalue()